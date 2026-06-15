"""
Генерация заявления студента:
  - DOCX для печати (студент ставит подпись)
  - PDF страница заявления
  - Объединённый PDF (заявление + сертификаты)
"""
from __future__ import annotations

import io
import os
from dataclasses import dataclass, field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_NAME = "Times New Roman"

_WIN_FONTS = "C:/Windows/Fonts"
_FONT_REGULAR = os.path.join(_WIN_FONTS, "times.ttf")
_FONT_BOLD = os.path.join(_WIN_FONTS, "timesbd.ttf")


@dataclass
class SubjectEntry:
    name: str = ""
    cert_sources: list[str] = field(default_factory=list)


@dataclass
class ApplicationPayload:
    student_name: str = ""
    faculty: str = ""
    department: str = ""
    group: str = ""
    course_year: str = ""
    phone: str = ""
    date: str = ""
    subjects: list[SubjectEntry] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _apply_font(run, *, bold: bool = False, italic: bool = False, size_pt: int = 12) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)


def _set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)


def _add_para(
    doc: Document,
    text: str = "",
    *,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    bold: bool = False,
    italic: bool = False,
    size_pt: int = 12,
    space_before_pt: float = 0,
    space_after_pt: float = 4,
    indent_cm: float = 0.0,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    if indent_cm:
        p.paragraph_format.first_line_indent = Cm(indent_cm)
    if text:
        run = p.add_run(text)
        _apply_font(run, bold=bold, italic=italic, size_pt=size_pt)


# ---------------------------------------------------------------------------
# DOCX generation
# ---------------------------------------------------------------------------
def build_application_docx(payload: ApplicationPayload) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    _set_default_font(doc)

    def right(text: str, size: int = 12, italic: bool = False, after: float = 2) -> None:
        _add_para(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  italic=italic, size_pt=size, space_after_pt=after)

    right("Председателю комиссии")
    right("Валиевой М.М.")
    right(f"от студента/магистранта {payload.course_year or '__'} курса")
    right(f"(группа) {payload.group or '___________'}")
    right(payload.student_name or "_________________________")
    right("ФИО", size=10, italic=True)
    right(payload.phone or "_________________________")
    right("телефон", size=10, italic=True, after=12)

    # Title
    _add_para(doc, "Заявление",
              align=WD_ALIGN_PARAGRAPH.CENTER,
              bold=True, size_pt=14,
              space_before_pt=6, space_after_pt=10)

    # Body
    _add_para(
        doc,
        "Прошу Вас признать результаты неформального обучения, по дисциплинам:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        indent_cm=1.25, space_after_pt=4,
    )

    for i in range(max(5, len(payload.subjects))):
        subj = payload.subjects[i].name if i < len(payload.subjects) else ""
        _add_para(doc, f"{i + 1}. {subj or ('_' * 55)}", space_after_pt=2)

    # Build cert list with sources
    if payload.subjects:
        cert_parts = []
        for s in payload.subjects:
            if s.name:
                if s.cert_sources:
                    cert_parts.append(f"{s.name} ({', '.join(s.cert_sources)})")
                else:
                    cert_parts.append(s.name)
        cert_list = "; ".join(cert_parts) if cert_parts else "_" * 60
    else:
        cert_list = "_" * 60
    _add_para(
        doc,
        "Документы, подтверждающие приобретение соответствующих результатов "
        f"обучения прилагаются: {cert_list}",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        indent_cm=1.25, space_before_pt=8, space_after_pt=6,
    )

    _add_para(
        doc,
        "Согласен на использование личных данных.",
        indent_cm=1.25, space_after_pt=24,
    )

    # Signature table (2 columns: empty | signature)
    tbl = doc.add_table(rows=3, cols=2)
    tbl.autofit = False
    tbl.columns[0].width = Cm(9)
    tbl.columns[1].width = Cm(7.5)

    def _tcell(cell, text: str, italic: bool = False, size: int = 12) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        _apply_font(run, italic=italic, size_pt=size)

    _tcell(tbl.rows[0].cells[0], "")
    _tcell(tbl.rows[0].cells[1], "_________________________")
    _tcell(tbl.rows[1].cells[1], "Подпись", italic=True, size=10)
    _tcell(tbl.rows[2].cells[1], payload.date or "«__» __________ 20__ г.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF generation (application page)
# ---------------------------------------------------------------------------
def build_application_pdf_page(payload: ApplicationPayload) -> bytes:
    import fitz  # PyMuPDF

    doc = fitz.open()
    page = doc.new_page(width=595.276, height=841.89)  # A4

    subjects_html = "".join(
        f"<p style='margin:2pt 0;'>{i + 1}. "
        f"{payload.subjects[i].name if i < len(payload.subjects) else '_' * 40}</p>"
        for i in range(max(5, len(payload.subjects)))
    )

    if payload.subjects:
        cert_parts = []
        for s in payload.subjects:
            if s.name:
                if s.cert_sources:
                    cert_parts.append(f"{s.name} ({', '.join(s.cert_sources)})")
                else:
                    cert_parts.append(s.name)
        cert_list = "; ".join(cert_parts) if cert_parts else "_" * 60
    else:
        cert_list = "_" * 60
    date_str = payload.date or "«__» __________ 20__ г."

    html = (
        "<div style='font-family:\"Times New Roman\",Times,serif;"
        "font-size:11pt;color:#000;'>"
        "<div style='text-align:right;line-height:1.7;margin-bottom:14pt;'>"
        "Председателю комиссии<br/>"
        "Валиевой М.М.<br/>"
        f"от студента/магистранта {payload.course_year or '__'} курса<br/>"
        f"(группа) {payload.group or '___________'}<br/>"
        f"{payload.student_name or '_________________________'}<br/>"
        "<span style='font-size:9pt;font-style:italic;'>ФИО</span><br/>"
        f"{payload.phone or '_________________________'}<br/>"
        "<span style='font-size:9pt;font-style:italic;'>телефон</span>"
        "</div>"
        "<p style='text-align:center;font-weight:bold;font-size:14pt;"
        "margin:8pt 0;'>Заявление</p>"
        "<p style='text-align:justify;margin:6pt 0;'>"
        "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;"
        "Прошу Вас признать результаты неформального обучения, по дисциплинам:"
        "</p>"
        + subjects_html
        + "<p style='text-align:justify;margin:10pt 0;'>"
        "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;"
        "Документы, подтверждающие приобретение соответствующих результатов "
        f"обучения прилагаются: {cert_list}"
        "</p>"
        "<p style='margin:6pt 0 20pt 40pt;'>"
        "Согласен на использование личных данных.</p>"
        "<div style='text-align:right;margin-top:16pt;line-height:1.8;'>"
        "_________________________<br/>"
        "<span style='font-size:9pt;font-style:italic;'>Подпись</span><br/>"
        f"{date_str}<br/>"
        "<span style='font-size:9pt;font-style:italic;'>Дата заявления</span>"
        "</div></div>"
    )

    rect = fitz.Rect(57, 57, 538, 784)
    if os.path.isdir(_WIN_FONTS):
        try:
            arch = fitz.Archive(_WIN_FONTS)
            page.insert_htmlbox(rect, html, archive=arch)
        except Exception:
            page.insert_htmlbox(rect, html)
    else:
        page.insert_htmlbox(rect, html)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Analysis summary page
# ---------------------------------------------------------------------------
def build_subject_analysis_page(
    payload: ApplicationPayload,
    cert_groups: list[list[bytes]],
    cert_group_names: list[list[str]] | None = None,
) -> bytes:
    """Страница 2: чёткая таблица — какой сертификат (Прил. №N, стр. X) к какой дисциплине."""
    import fitz

    # Вычисляем номера вложений и страниц
    page_cursor = 3  # стр.1=заявление, стр.2=эта страница, стр.3+=сертификаты
    attach_num = 1
    total_attachments = 0
    rows_html = ""

    for si, subj in enumerate(payload.subjects):
        group = cert_groups[si] if si < len(cert_groups) else []
        names = cert_group_names[si] if cert_group_names and si < len(cert_group_names) else []

        # Заголовок дисциплины
        rows_html += (
            f"<tr style='background:#e8f0fb;'>"
            f"<td colspan='3' style='padding:6pt 8pt;font-weight:700;font-size:10.5pt;"
            f"border-top:2px solid #c7d9f5;'>{si + 1}.&nbsp;{subj.name}</td></tr>"
        )

        if group:
            for ci, cert_bytes in enumerate(group):
                src = subj.cert_sources[ci] if ci < len(subj.cert_sources) else ""
                fname = names[ci] if ci < len(names) else ""
                src_short = (src[:75] + "…") if len(src) > 75 else src

                # Считаем страницы сертификата для точной ссылки
                n_pages = 1
                try:
                    pdf_data = cert_data_to_pdf(cert_bytes, fname)
                    c_doc = fitz.open("pdf", pdf_data)
                    n_pages = len(c_doc)
                    c_doc.close()
                except Exception:
                    pass

                page_ref = (
                    f"стр.&nbsp;{page_cursor}"
                    if n_pages == 1
                    else f"стр.&nbsp;{page_cursor}–{page_cursor + n_pages - 1}"
                )

                rows_html += (
                    f"<tr><td style='padding:5pt 8pt;color:#1d4ed8;font-weight:700;"
                    f"white-space:nowrap;'>Прил.&nbsp;№{attach_num}</td>"
                    f"<td style='padding:5pt 8pt;color:#475569;font-size:8.5pt;"
                    f"white-space:nowrap;'>{page_ref}</td>"
                    f"<td style='padding:5pt 8pt;font-size:8.5pt;'>{src_short or '—'}</td></tr>"
                )
                page_cursor += n_pages
                attach_num += 1
                total_attachments += 1
        else:
            # Нет прикреплённых файлов
            src_hint = "; ".join(subj.cert_sources[:3]) if subj.cert_sources else ""
            rows_html += (
                f"<tr><td colspan='3' style='padding:5pt 8pt;color:#94a3b8;"
                f"font-size:8.5pt;font-style:italic;'>файл не прикреплён"
                f"{(' — ' + src_hint) if src_hint else ''}</td></tr>"
            )

    html = (
        "<div style='font-family:\"Times New Roman\",Times,serif;font-size:10pt;color:#000;'>"
        "<h2 style='text-align:center;font-size:12pt;margin:0 0 6pt;'>Приложение к заявлению</h2>"
        "<p style='text-align:center;font-size:9pt;color:#475569;margin:0 0 10pt;'>"
        "Распределение приложенных сертификатов по дисциплинам</p>"
        f"<p style='margin:0 0 10pt;'>"
        f"<b>Студент:</b> {payload.student_name or '—'}&nbsp;&nbsp;"
        f"<b>Группа:</b> {payload.group or '—'}&nbsp;&nbsp;"
        f"<b>Курс:</b> {payload.course_year or '—'}</p>"
        "<table style='width:100%;border-collapse:collapse;border:1px solid #c7d9f5;'>"
        "<thead><tr style='background:#1e3a5f;color:#fff;'>"
        "<th style='padding:6pt 8pt;text-align:left;font-size:9pt;width:15%;'>Вложение</th>"
        "<th style='padding:6pt 8pt;text-align:left;font-size:9pt;width:16%;'>Страница</th>"
        "<th style='padding:6pt 8pt;text-align:left;font-size:9pt;'>Источник / URL сертификата</th>"
        "</tr></thead>"
        f"<tbody>{rows_html}</tbody></table>"
        f"<p style='font-size:8pt;color:#475569;margin-top:8pt;'>"
        f"<b>Итого вложений: {total_attachments}.</b>&nbsp;"
        "Сертификаты расположены в документе сразу после данной страницы "
        "(Прил.&nbsp;№1&nbsp;=&nbsp;стр.&nbsp;3).</p>"
        "</div>"
    )

    doc = fitz.open()
    page = doc.new_page(width=595.276, height=841.89)
    rect = fitz.Rect(40, 40, 555, 800)
    if os.path.isdir(_WIN_FONTS):
        try:
            arch = fitz.Archive(_WIN_FONTS)
            page.insert_htmlbox(rect, html, archive=arch)
        except Exception:
            page.insert_htmlbox(rect, html)
    else:
        page.insert_htmlbox(rect, html)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Image → PDF helper
# ---------------------------------------------------------------------------
_IMAGE_EXTS = {"jpg", "jpeg", "png", "webp", "gif", "bmp", "tiff", "tif"}


def image_to_pdf_bytes(image_bytes: bytes, filename: str = "") -> bytes:
    """Convert an image file to a one-page PDF."""
    import fitz
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "jpg"
    if ext not in _IMAGE_EXTS:
        ext = "jpg"
    img_doc = fitz.open(stream=image_bytes, filetype=ext)
    pdf_bytes = img_doc.convert_to_pdf()
    img_doc.close()
    return pdf_bytes


def cert_data_to_pdf(data: bytes, filename: str = "") -> bytes:
    """Return PDF bytes from a certificate — converts images automatically."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in _IMAGE_EXTS:
        return image_to_pdf_bytes(data, filename)
    # Try to open as PDF; if that fails, try common image types
    import fitz
    try:
        fitz.open("pdf", data).close()
        return data
    except Exception:
        pass
    for try_ext in ("png", "jpg", "webp"):
        try:
            img_doc = fitz.open(stream=data, filetype=try_ext)
            pdf_bytes = img_doc.convert_to_pdf()
            img_doc.close()
            return pdf_bytes
        except Exception:
            continue
    return data  # fallback: return as-is


# ---------------------------------------------------------------------------
# PDF merge
# ---------------------------------------------------------------------------
def build_combined_pdf(
    app_pdf_bytes: bytes,
    cert_pdfs: list[bytes],
    payload: ApplicationPayload | None = None,
    cert_groups: list[list[bytes]] | None = None,
    cert_filenames: list[str] | None = None,
) -> bytes:
    import fitz

    merged = fitz.open()

    # Page 1: application
    app_doc = fitz.open("pdf", app_pdf_bytes)
    merged.insert_pdf(app_doc)

    # Page 2: subject analysis (if payload provided)
    if payload and payload.subjects:
        groups = cert_groups or [[] for _ in payload.subjects]
        # Build per-group filenames from flat cert_filenames
        group_names: list[list[str]] | None = None
        if cert_filenames and cert_groups:
            group_names = []
            fi = 0
            for g in cert_groups:
                gn: list[str] = []
                for _ in g:
                    gn.append(cert_filenames[fi] if fi < len(cert_filenames) else "")
                    fi += 1
                group_names.append(gn)
        try:
            analysis_bytes = build_subject_analysis_page(payload, groups, group_names)
            analysis_doc = fitz.open("pdf", analysis_bytes)
            merged.insert_pdf(analysis_doc)
        except Exception:
            pass

    # Remaining pages: certificates (auto-convert images to PDF)
    for i, cert_bytes in enumerate(cert_pdfs):
        if not cert_bytes:
            continue
        try:
            fname = cert_filenames[i] if cert_filenames and i < len(cert_filenames) else ""
            pdf_data = cert_data_to_pdf(cert_bytes, fname)
            cert_doc = fitz.open("pdf", pdf_data)
            merged.insert_pdf(cert_doc)
        except Exception:
            pass

    buf = io.BytesIO()
    merged.save(buf)
    return buf.getvalue()
