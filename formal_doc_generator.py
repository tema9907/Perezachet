"""Генерация DOCX для модуля формального перезачёта: Экспертное заключение + Протокол.

Повторяет структуру образцов из папки «формальное/»:
  - Заключение (альбомная): шапка-реквизиты + таблица (9 колонок) + подписи по кафедрам.
  - Протокол (книжная): двуязычный бланк + состав комиссии + СЛУШАЛИ/ПОСТАНОВИЛИ +
    та же таблица + подписи.

Переиспользует низкоуровневые helper'ы из doc_generator (шрифт, ячейки, границы).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import doc_generator as dg


# ---------------------------------------------------------------------------
# Payload
# ---------------------------------------------------------------------------
@dataclass
class FormalRow:
    plan_name: str = ""          # дисциплина по учебному плану СКУ
    plan_credits: str = ""       # кредиты по плану
    fact_name: str = ""          # дисциплина формального обучения (из транскрипта)
    fact_credits: str = ""       # кредиты/часы факт
    fact_grade: str = ""         # баллы(оценка), напр. «68(C)»
    compliance: str = ""         # полное / частичное / не соответствует
    note: str = ""               # примечание
    final_grade: str = ""        # зачтённая оценка, напр. «68(C)» или «АР»


@dataclass
class Kafedra:
    name: str
    experts: list[str] = field(default_factory=lambda: ["", "", ""])


# Кафедры и эксперты по умолчанию — как в образце заключения.
DEFAULT_KAFEDRAS: list[Kafedra] = [
    Kafedra("Кафедра ИКТ", [
        "Борамбаева Г.М., сеньор-лектор кафедры ИКТ",
        "Курмашева Л.Б., сеньор-лектор кафедры ИКТ",
        "Семейников А.Н., ассистент кафедры ИКТ",
    ]),
    Kafedra("Кафедра Иностранные языки"),
    Kafedra("Кафедра Математика и физика"),
    Kafedra("Кафедра Физическая культура и спорт"),
    Kafedra("Кафедра История Казахстана и право"),
    Kafedra("Кафедра Педагогика и психология"),
    Kafedra("Кафедра Экономика, учет и финансы"),
    Kafedra("Кафедра Практический казахский язык"),
]

# Состав комиссии по умолчанию — как в образце протокола.
DEFAULT_CHAIR = "Валиева М.М. – директор ДАД"
DEFAULT_DEPUTY = "Иванцова О.А. – руководитель ОР"
DEFAULT_MEMBERS = [
    "Икласова К.Е. – декан факультета Инженерии и цифровых технологий;",
    "Курмашев И.Г. – заведующий кафедрой «Информационно-коммуникационные технологии»;",
    "Савинкин В.В. – заведующий кафедрой «Транспорт и машиностроение»;",
    "Кашевкин А.А. – заведующий кафедрой «Энергетика и радиоэлектроника»;",
    "Аубакирова Б.Б. – заведующий кафедрой «Строительство и дизайн»;",
]
DEFAULT_SECRETARY = "Нурпеисов Т.К. – лаборант ФИЦТ."
DEFAULT_REPORTER = "Курмашев И.Г., заведующий кафедрой «Информационно-коммуникационные технологии»"
# Подписи внизу протокола: (должность, ФИО)
DEFAULT_SIGN_MEMBERS = [
    ("декан ФИЦТ", "Икласова К.Е."),
    ("к.т.н., доцент кафедры ИКТ", "Курмашев И.Г."),
    ("к.т.н., доцент кафедры ТиМ", "Савинкин В.В."),
    ("PhD, доцент кафедры ЭиР", "Кашевкин А.А."),
    ("PhD, доцент кафедры СиД", "Аубакирова Б.Б."),
]


@dataclass
class FormalPayload:
    student_name: str = ""
    target_program: str = ""        # «6В06104 «Вычислительная техника и ПО»»
    course: str = ""
    source_institution: str = ""    # «НАО «СКУ им. М.Козыбаева»»
    source_program: str = ""        # «6В06111 «Информационные системы в управлении»»
    prev_doc_type: str = "транскрипт"
    rows: list[FormalRow] = field(default_factory=list)

    # реквизиты протокола
    protocol_number: str = ""
    protocol_date: str = ""         # «24.08.2026»
    group: str = ""                 # «ВТиПО-25к»
    enroll_semester: str = ""       # «3»
    academic_difference: str = ""   # «11» — если пусто, считается автоматически

    # состав
    kafedras: list[Kafedra] = field(default_factory=lambda: [Kafedra(k.name, list(k.experts)) for k in DEFAULT_KAFEDRAS])
    chair: str = DEFAULT_CHAIR
    deputy: str = DEFAULT_DEPUTY
    members: list[str] = field(default_factory=lambda: list(DEFAULT_MEMBERS))
    secretary: str = DEFAULT_SECRETARY
    reporter: str = DEFAULT_REPORTER
    sign_members: list[tuple[str, str]] = field(default_factory=lambda: list(DEFAULT_SIGN_MEMBERS))


# ---------------------------------------------------------------------------
# Русские числительные (для «__11__ (одиннадцать) кредитов»)
# ---------------------------------------------------------------------------
_ONES = ["ноль", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь",
         "девять", "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
         "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят",
         "семьдесят", "восемьдесят", "девяносто"]


def number_to_words(n: int) -> str:
    if n < 0:
        return str(n)
    if n < 20:
        return _ONES[n]
    if n < 100:
        t, o = divmod(n, 10)
        return _TENS[t] + ((" " + _ONES[o]) if o else "")
    return str(n)


def compute_academic_difference(rows: list[FormalRow]) -> int:
    """Сумма кредитов дисциплин, помеченных как академ. разница (АР / нет соответствия)."""
    total = 0
    for r in rows:
        is_ar = (r.final_grade.strip().upper() == "АР") or not r.fact_name.strip()
        if is_ar:
            try:
                total += int(float((r.plan_credits or "0").replace(",", ".")))
            except (ValueError, TypeError):
                pass
    return total


# ---------------------------------------------------------------------------
# Общая таблица дисциплин (9 колонок) — используется и в заключении, и в протоколе
# ---------------------------------------------------------------------------
_H_ROW0 = {
    0: "№",
    1: "По учебному плану СКУ им.М.Козыбаева",       # 1-2
    3: "Фактически изучено претендентом и рассмотрено экспертной группой",  # 3-5
    6: "Заключение экспертной группы",               # 6-7
    8: "Зачтенная оценка по дисциплине учебного плана",
}
_H_ROW1 = {
    1: "Наименование перезачитываемых дисциплин и видов учебной работы",
    2: "Количество кредитов",
    3: "Наименование дисциплины формального обучения",
    4: "Количество кредитов (часов)",
    5: "Баллы (оценка)",
    6: "Соответствие целей программы обучения, объёма программы и оценки "
       "(полное, частичное, не соответствует)",
    7: "Примечание",
}
_CENTER_COLS = {0, 2, 4, 5, 6, 8}


def _build_disc_table(doc, rows: list[FormalRow], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=2 + len(rows), cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # merges (сначала мерджим, потом пишем — иначе Word дублирует текст)
    table.rows[0].cells[1].merge(table.rows[0].cells[2])
    table.rows[0].cells[3].merge(table.rows[0].cells[5])
    table.rows[0].cells[6].merge(table.rows[0].cells[7])
    for col in (0, 8):
        table.rows[0].cells[col].merge(table.rows[1].cells[col])

    for ci, text in _H_ROW0.items():
        dg._set_cell_text(table.rows[0].cells[ci], text, bold=True, align_center=True)
    for ci, text in _H_ROW1.items():
        dg._set_cell_text(table.rows[1].cells[ci], text, bold=True, align_center=True)

    for ri, r in enumerate(rows, start=2):
        values = [
            str(ri - 1), r.plan_name, r.plan_credits, r.fact_name,
            r.fact_credits, r.fact_grade, r.compliance, r.note, r.final_grade,
        ]
        cells = table.rows[ri].cells
        for ci, v in enumerate(values):
            dg._set_cell_text(cells[ci], v, align_center=(ci in _CENTER_COLS))

    dg._set_col_widths(table, widths_cm)
    dg._add_borders(table)
    dg._set_cell_margins(table)


def _p(doc, text: str = "", *, bold: bool = False, center: bool = False,
       italic: bool = False, size: int = 12):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    dg._apply_font(run, bold=bold, size_pt=size)
    run.italic = italic
    return para


def _apply_default_font(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = dg.FONT_NAME
    style.font.size = Pt(12)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), dg.FONT_NAME)


# ---------------------------------------------------------------------------
# Заключение (альбомная)
# ---------------------------------------------------------------------------
CONCLUSION_WIDTHS = [0.8, 4.8, 2.0, 4.8, 2.2, 1.8, 4.0, 1.8, 2.2]  # сумма ≈ 24.4


def build_conclusion(payload: FormalPayload) -> bytes:
    doc = Document()
    dg._set_landscape(doc.sections[0])
    _apply_default_font(doc)

    _p(doc, "Экспертное заключение", bold=True, center=True)
    _p(doc, "по определению соответствия результатов формального обучения "
            "по профилю подготовки", bold=True, center=True)
    _p(doc, "")

    name = dg.format_student_name(payload.student_name).replace("\n", " ")
    _p(doc, f"ФИО обучающегося    {name}")
    _p(doc, f"Код и наименование ОП    {payload.target_program}")
    _p(doc, f"Курс    {payload.course}")
    prev = f"{payload.prev_doc_type}, {payload.source_institution}, специальность {payload.source_program}"
    _p(doc, f"Документ предыдущего образования    {prev}")
    _p(doc, "наименование, номер, дата предоставленного документа", italic=True, size=10)
    _p(doc, "")

    _build_disc_table(doc, payload.rows or [FormalRow()], CONCLUSION_WIDTHS)

    # подписи по кафедрам
    _p(doc, "")
    nbsp = " "
    for kaf in payload.kafedras:
        _p(doc, kaf.name, bold=True)
        _p(doc, "Эксперты:", bold=True)
        for expert in (kaf.experts or ["", "", ""]):
            label = expert.strip() or ("_" * 38)
            line = doc.add_paragraph()
            line.paragraph_format.space_after = Pt(2)
            run = line.add_run(f"{label}\t{nbsp*5}______________{nbsp*4}______________")
            dg._apply_font(run, bold=False, size_pt=12)
            sub = doc.add_paragraph()
            sub.paragraph_format.space_after = Pt(4)
            srun = sub.add_run(f"\t\t\t(ФИО, должность)\t\t{nbsp*8}(подпись)\t{nbsp*12}(дата)")
            dg._apply_font(srun, bold=False, size_pt=9)
            srun.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Протокол (книжная)
# ---------------------------------------------------------------------------
PROTOCOL_WIDTHS = [0.8, 2.7, 1.6, 3.6, 1.8, 1.4, 3.2, 1.3, 1.7]  # сумма ≈ 18.1


def _protocol_header_table(doc, number: str, date: str) -> None:
    table = doc.add_table(rows=4, cols=3)
    table.autofit = True
    kz = [
        "«МАНАШ ҚОЗЫБАЕВ АТЫНДАҒЫ СОЛТҮСТІК ҚАЗАҚСТАН УНИВЕРСИТЕТІ» "
        "коммерциялық емес акционерлік қоғамы",
        "ХАТТАМА КӨШІРМЕСІ",
        f"№ {number}   {date} ж.",
        "Петропавл қ.",
    ]
    ru = [
        "Некоммерческое акционерное общество «СЕВЕРО-КАЗАХСТАНСКИЙ УНИВЕРСИТЕТ "
        "ИМЕНИ МАНАША КОЗЫБАЕВА»",
        "ВЫПИСКА ИЗ ПРОТОКОЛА",
        f"№ {number}   {date} г.",
        "г. Петропавловск",
    ]
    for i in range(4):
        dg._set_cell_text(table.rows[i].cells[0], kz[i], bold=(i in (1,)), align_center=True, size_pt=10)
        dg._set_cell_text(table.rows[i].cells[1], "", align_center=True)
        dg._set_cell_text(table.rows[i].cells[2], ru[i], bold=(i in (1,)), align_center=True, size_pt=10)
    dg._set_col_widths(table, [8.5, 1.0, 8.5])


def build_protocol(payload: FormalPayload) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.0)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    _apply_default_font(doc)

    _protocol_header_table(doc, payload.protocol_number, payload.protocol_date)
    _p(doc, "заседания Комиссии по признанию результатов", center=True)
    _p(doc, "формального и неформального образования", center=True)
    _p(doc, "")
    _p(doc, f"Председатель: {payload.chair}")
    _p(doc, f"Заместитель председателя: {payload.deputy}")
    _p(doc, "Члены комиссии:")
    for i, mem in enumerate(payload.members, start=1):
        _p(doc, f"{i}) {mem}")
    _p(doc, f"Секретарь: {payload.secretary}")
    _p(doc, "")
    _p(doc, "ПОВЕСТКА ДНЯ:", bold=True)
    _p(doc, f"1. О признании результатов формального обучения в НАО «СКУ им. М.Козыбаева». "
            f"Докладчик: {payload.reporter}.")
    _p(doc, "1. СЛУШАЛИ:", bold=True)
    reporter_short = payload.reporter.split(",")[0].strip()
    _p(doc, f"{reporter_short} предложил признать результаты формального обучения в НАО "
            f"«СКУ им. М. Козыбаева» {payload.student_name} по образовательной программе "
            f"{payload.target_program} на основании экспертного заключения (по определению "
            f"соответствия результатов формального обучения по профилю подготовки, по "
            f"определению соответствия результатов формального обучения по дисциплинам "
            f"цикла ООД и БД).")
    _p(doc, f"– из {payload.source_institution}, {payload.source_program}.")
    _p(doc, "")

    _build_disc_table(doc, payload.rows or [FormalRow()], PROTOCOL_WIDTHS)

    # ПОСТАНОВИЛИ
    diff = payload.academic_difference.strip()
    if not diff:
        diff = str(compute_academic_difference(payload.rows))
    try:
        diff_words = number_to_words(int(diff))
    except (ValueError, TypeError):
        diff_words = ""
    _p(doc, "")
    _p(doc, "ПОСТАНОВИЛИ:", bold=True)
    _p(doc, f"1.1 Признать результаты формального обучения в НАО «СКУ им. М. Козыбаева» "
            f"{payload.student_name} для обучения по ОП {payload.target_program}.")
    _p(doc, f"1.2 Рекомендовать зачислить {payload.student_name} в группу "
            f"{payload.group} на {payload.enroll_semester} семестр.")
    _p(doc, f"1.3. Академическая разница составляет __{diff}__ ({diff_words}) кредитов.")
    _p(doc, "")

    # подписи
    nbsp = " "
    _p(doc, f"Председатель:{nbsp*10}{payload.chair.split('–')[0].strip()}")
    _p(doc, f"Зам.председателя:{nbsp*6}{payload.deputy.split('–')[0].strip()}")
    _p(doc, "Члены комиссии:")
    for role, fio in payload.sign_members:
        _p(doc, f"{role}{nbsp*10}{fio}")
    _p(doc, f"Секретарь{nbsp*10}{payload.secretary.split('–')[0].strip()}")
    _p(doc, "")
    _p(doc, f"С протоколом ознакомлен(а)  {payload.student_name}{nbsp*6}______________")
    _p(doc, "(ФИО обучающегося)                                        (подпись)",
       italic=True, size=9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
